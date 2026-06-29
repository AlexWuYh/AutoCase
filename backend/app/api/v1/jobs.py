"""Generation jobs API endpoints."""
from __future__ import annotations

from io import BytesIO, StringIO
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ...database import get_db
from ...deps import get_current_user
from ...models.case import GeneratedCase
from ...models.job import GenerationJob
from ...models.requirement import Requirement
from ...models.requirement_group import RequirementGroup
from ...models.user import User, UserRole
from ...schemas.common import Message, PageResponse
from ...schemas.job import GeneratedCaseOut, JobCreate, JobOut
from ...tasks.generation import generate_job

router = APIRouter()


def _job_out(job: GenerationJob) -> JobOut:
    total = job.total or 0
    finished = job.finished or 0
    return JobOut(
        id=job.id,
        group_id=job.group_id,
        user_id=job.user_id,
        status=job.status,
        total=total,
        finished=finished,
        percent=round((finished / total * 100) if total > 0 else 0, 1),
        error=job.error,
        group_name=job.group.name if job.group else None,
        user_name=job.user.username if job.user else None,
        llm_config_name=job.llm_config.name if job.llm_config else None,
        prompt_name=job.prompt.name if job.prompt else None,
        created_at=job.created_at,
        started_at=job.started_at,
        finished_at=job.finished_at,
    )


def _can_access(job: GenerationJob | None, user: User) -> bool:
    if job is None:
        return False
    return user.role == UserRole.ADMIN or job.user_id == user.id


# ─── CRUD ───────────────────────────────────────────────────────────────

@router.get("", response_model=PageResponse[JobOut], tags=["jobs"])
def list_jobs(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    status_filter: str | None = Query(None, alias="status"),
    group_id: int | None = Query(None),
) -> PageResponse[JobOut]:
    stmt = select(GenerationJob)
    if current_user.role != UserRole.ADMIN:
        stmt = stmt.where(GenerationJob.user_id == current_user.id)
    if status_filter:
        stmt = stmt.where(GenerationJob.status == status_filter)
    if group_id is not None:
        stmt = stmt.where(GenerationJob.group_id == group_id)
    stmt = stmt.order_by(GenerationJob.id.desc())

    count_stmt = select(func.count()).select_from(stmt.subquery())
    total = db.execute(count_stmt).scalar_one()
    rows = db.execute(stmt.offset((page - 1) * page_size).limit(page_size)).scalars().all()
    return PageResponse(items=[_job_out(j) for j in rows], total=total, page=page, page_size=page_size)


@router.post("", response_model=JobOut, status_code=status.HTTP_202_ACCEPTED, tags=["jobs"])
def create_job(
    payload: JobCreate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> JobOut:
    group = db.get(RequirementGroup, payload.group_id)
    if group is None:
        raise HTTPException(status_code=404, detail="需求集不存在")
    if current_user.role != UserRole.ADMIN and group.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="无权操作此需求集")
    count = db.query(Requirement).filter(Requirement.group_id == group.id).count()
    if count == 0:
        raise HTTPException(status_code=400, detail="需求集中没有功能点，请先添加")

    job = GenerationJob(
        group_id=group.id,
        user_id=current_user.id,
        llm_config_id=payload.llm_config_id,
        prompt_id=payload.prompt_id,
        status="pending",
        total=count,
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    generate_job.delay(job.id)
    return _job_out(job)


@router.get("/{job_id}", response_model=JobOut, tags=["jobs"])
def get_job(job_id: int, current_user: Annotated[User, Depends(get_current_user)], db: Annotated[Session, Depends(get_db)]) -> JobOut:
    job = db.get(GenerationJob, job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="任务不存在")
    if not _can_access(job, current_user):
        raise HTTPException(status_code=403, detail="无权访问")
    return _job_out(job)


@router.delete("/{job_id}", response_model=Message, tags=["jobs"])
def cancel_job(job_id: int, current_user: Annotated[User, Depends(get_current_user)], db: Annotated[Session, Depends(get_db)]) -> Message:
    job = db.get(GenerationJob, job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="任务不存在")
    if not _can_access(job, current_user):
        raise HTTPException(status_code=403, detail="无权操作")
    if job.status not in ("pending", "running"):
        raise HTTPException(status_code=400, detail=f"任务状态为 {job.status}，无法取消")
    job.status = "canceled"
    db.commit()
    return Message(message="已取消")


@router.get("/{job_id}/cases", response_model=PageResponse[GeneratedCaseOut], tags=["jobs"])
def list_cases(
    job_id: int, current_user: Annotated[User, Depends(get_current_user)], db: Annotated[Session, Depends(get_db)],
    page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=500),
) -> PageResponse[GeneratedCaseOut]:
    job = db.get(GenerationJob, job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="任务不存在")
    if not _can_access(job, current_user):
        raise HTTPException(status_code=403, detail="无权访问")
    stmt = select(GeneratedCase).where(GeneratedCase.job_id == job.id).order_by(GeneratedCase.id)
    total = len(db.execute(stmt).all())
    rows = db.execute(stmt.offset((page - 1) * page_size).limit(page_size)).scalars().all()
    return PageResponse(items=[GeneratedCaseOut.model_validate(r) for r in rows], total=total, page=page, page_size=page_size)


@router.get("/{job_id}/export", tags=["jobs"])
def export_job(
    job_id: int, current_user: Annotated[User, Depends(get_current_user)], db: Annotated[Session, Depends(get_db)],
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
) -> StreamingResponse:
    job = db.get(GenerationJob, job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="任务不存在")
    if not _can_access(job, current_user):
        raise HTTPException(status_code=403, detail="无权访问")
    cases = db.query(GeneratedCase).filter(GeneratedCase.job_id == job.id).order_by(GeneratedCase.id).all()

    if format == "csv":
        import csv
        output = StringIO()
        w = csv.writer(output)
        w.writerow(["用例ID", "所属模块", "用例名称", "前置条件", "步骤", "预期", "关键词", "优先级", "用例类型", "适用阶段"])
        for c in cases:
            w.writerow([c.case_id, c.module, c.name, c.preconditions,
                        "\n".join(f"{i+1}. {s}" for i, s in enumerate(c.steps)),
                        "\n".join(f"{i+1}. {e}" for i, e in enumerate(c.expected)),
                        c.keywords, c.priority, c.case_type, c.stage])
        output.seek(0)
        return StreamingResponse(output, media_type="text/csv", headers={"Content-Disposition": "attachment; filename=cases.csv"})

    wb = Workbook()
    ws = wb.active
    ws.title = "TestCases"
    headers = ["用例ID", "所属模块", "用例名称", "前置条件", "步骤", "预期", "关键词", "优先级", "用例类型", "适用阶段"]
    ws.append(headers)
    for c in cases:
        ws.append([c.case_id, c.module, c.name, c.preconditions,
                   "\n".join(f"{i+1}. {s}" for i, s in enumerate(c.steps)),
                   "\n".join(f"{i+1}. {e}" for i, e in enumerate(c.expected)),
                   c.keywords, c.priority, c.case_type, c.stage])

    header_fill = PatternFill("solid", fgColor="E8EEF7")
    header_font = Font(bold=True)
    thin = Side(border_style="thin", color="CBD5E1")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.alignment = wrap
            cell.border = bdr
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = header_font
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        mx = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(mx + 2, 60)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=job_{job_id}_testcases.xlsx"})
